"""Tulis hasil riset ke Excel (3 sheet) atau 3 file CSV."""

from __future__ import annotations

import csv
from dataclasses import fields
from pathlib import Path

from openpyxl import Workbook
from openpyxl.utils import get_column_letter

from .models import Produk, Profil, Video


def _tabel(model: type, baris: list) -> tuple[list[str], list[list]]:
    kolom = [f.name for f in fields(model)]
    return kolom, [[obj.ke_dict()[k] for k in kolom] for obj in baris]


def _semua_tabel(profil: Profil, videos: list[Video], produk: list[Produk]) -> dict:
    return {
        "profil": _tabel(Profil, [profil]),
        "video": _tabel(Video, videos),
        "produk": _tabel(Produk, produk),
    }


def tulis_excel(path: Path, profil: Profil, videos: list[Video], produk: list[Produk]) -> Path:
    wb = Workbook()
    wb.remove(wb.active)
    for nama, (kolom, baris) in _semua_tabel(profil, videos, produk).items():
        ws = wb.create_sheet(nama)
        ws.append(kolom)
        for b in baris:
            ws.append(b)
        ws.freeze_panes = "A2"
        for i, k in enumerate(kolom, start=1):
            lebar = max([len(k)] + [len(str(b[i - 1] or "")) for b in baris])
            ws.column_dimensions[get_column_letter(i)].width = min(lebar + 2, 60)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def tulis_csv(
    folder: Path, awalan: str, profil: Profil, videos: list[Video], produk: list[Produk]
) -> list[Path]:
    folder.mkdir(parents=True, exist_ok=True)
    hasil = []
    for nama, (kolom, baris) in _semua_tabel(profil, videos, produk).items():
        path = folder / f"{awalan}_{nama}.csv"
        # utf-8-sig supaya Excel membaca huruf/emoji dengan benar
        with path.open("w", newline="", encoding="utf-8-sig") as f:
            tulis = csv.writer(f)
            tulis.writerow(kolom)
            tulis.writerows(baris)
        hasil.append(path)
    return hasil
