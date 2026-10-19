import json
from contextlib import nullcontext

from openpyxl import load_workbook

from riset_tiktok.browser import Diblokir, GagalAkses
from riset_tiktok.cli import main


class SumberPalsu:
    def __init__(self, fixture_teks, produk_diblokir=False, profil_diblokir=False):
        self.baca = fixture_teks
        self.produk_diblokir = produk_diblokir
        self.profil_diblokir = profil_diblokir
        self.dipanggil = []

    def riset_profil(self, username, maks_video):
        if self.profil_diblokir:
            raise Diblokir("captcha")
        self.dipanggil.append(("profil", username))
        data = json.loads(self.baca("item_list.json"))
        # halaman kedua mengulang video pertama: harus di-dedupe
        return self.baca("profil.html"), [data, {"itemList": data["itemList"][:1]}]

    def produk_html(self, product_id):
        self.dipanggil.append(("produk", product_id))
        if self.produk_diblokir:
            raise Diblokir("captcha")
        if product_id == "1729384756":
            return self.baca("produk.html"), f"https://shop/{product_id}"
        return "<html></html>", f"https://shop/{product_id}"


def _jalankan(argv, sumber):
    return main(argv, buat_sumber=lambda _args: nullcontext(sumber))


def test_cli_excel_lengkap(tmp_path, fixture_teks):
    sumber = SumberPalsu(fixture_teks)
    keluar = tmp_path / "hasil.xlsx"
    assert _jalankan(["@tokocontoh", "--keluar", str(keluar)], sumber) == 0
    assert sumber.dipanggil[0] == ("profil", "tokocontoh")

    wb = load_workbook(keluar)
    assert len(list(wb["video"].values)) == 4  # header + 3 video unik
    baris = list(wb["produk"].values)
    kolom = list(baris[0])
    produk = [dict(zip(kolom, b, strict=True)) for b in baris[1:]]
    assert [p["product_id"] for p in produk] == ["1729384756", "1729000001"]
    serum = produk[0]
    assert serum["nama"] == "Serum Wajah Glowing 30ml"
    assert serum["jumlah_video"] == 2
    assert serum["total_views_video"] == 170000
    assert produk[1]["nama"] in (None, "")  # halaman tanpa data -> detail kosong


def test_cli_maks_video_dan_maks_produk(tmp_path, fixture_teks):
    sumber = SumberPalsu(fixture_teks)
    keluar = tmp_path / "hasil.xlsx"
    assert _jalankan(["tokocontoh", "--maks-video", "1", "--maks-produk", "0",
                      "--keluar", str(keluar)], sumber) == 0
    wb = load_workbook(keluar)
    assert len(list(wb["video"].values)) == 2
    assert len(list(wb["produk"].values)) == 2
    assert not any(p[0] == "produk" for p in sumber.dipanggil)


def test_cli_csv(tmp_path, fixture_teks):
    assert _jalankan(["tokocontoh", "--csv", "--keluar", str(tmp_path)],
                     SumberPalsu(fixture_teks)) == 0
    assert (tmp_path / "hasil_tokocontoh_produk.csv").exists()


def test_cli_produk_diblokir_tetap_ekspor(tmp_path, fixture_teks):
    sumber = SumberPalsu(fixture_teks, produk_diblokir=True)
    keluar = tmp_path / "hasil.xlsx"
    assert _jalankan(["tokocontoh", "--keluar", str(keluar)], sumber) == 0
    # setelah blokir pertama, produk berikutnya tidak dibuka lagi
    assert [p for p in sumber.dipanggil if p[0] == "produk"] == [("produk", "1729384756")]
    assert len(list(load_workbook(keluar)["produk"].values)) == 3


def test_cli_profil_diblokir_gagal(tmp_path, fixture_teks, capsys):
    sumber = SumberPalsu(fixture_teks, profil_diblokir=True)
    assert _jalankan(["tokocontoh", "--keluar", str(tmp_path / "x.xlsx")], sumber) == 1
    assert "captcha" in capsys.readouterr().err
    assert not (tmp_path / "x.xlsx").exists()


def test_cli_username_kosong(fixture_teks):
    assert _jalankan(["@"], SumberPalsu(fixture_teks)) == 2


class SumberOffline(SumberPalsu):
    def riset_profil(self, username, maks_video):
        raise GagalAkses("tidak bisa membuka https://www.tiktok.com/@x: ERR_TUNNEL")


def test_cli_jaringan_gagal_pesan_jelas(tmp_path, fixture_teks, capsys):
    assert _jalankan(["x", "--keluar", str(tmp_path / "x.xlsx")],
                     SumberOffline(fixture_teks)) == 1
    assert "Periksa koneksi" in capsys.readouterr().err


class SumberProdukGagal(SumberPalsu):
    def produk_html(self, product_id):
        self.dipanggil.append(("produk", product_id))
        raise GagalAkses("timeout")


def test_cli_produk_gagal_dilewati_lanjut(tmp_path, fixture_teks):
    sumber = SumberProdukGagal(fixture_teks)
    assert _jalankan(["tokocontoh", "--keluar", str(tmp_path / "x.xlsx")], sumber) == 0
    assert len([p for p in sumber.dipanggil if p[0] == "produk"]) == 2
