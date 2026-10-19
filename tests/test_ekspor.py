import csv

from openpyxl import load_workbook

from riset_tiktok.ekspor import tulis_csv, tulis_excel
from riset_tiktok.models import Produk, Profil, Video


def _data():
    profil = Profil(username="tokocontoh", followers=10)
    videos = [Video(video_id="1", views=100, product_ids=["9", "8"]), Video(video_id="2")]
    produk = [Produk(product_id="9", nama="Serum 😊", harga=89000.0, jumlah_video=1)]
    return profil, videos, produk


def test_tulis_excel_tiga_sheet(tmp_path):
    path = tulis_excel(tmp_path / "hasil.xlsx", *_data())
    wb = load_workbook(path)
    assert wb.sheetnames == ["profil", "video", "produk"]
    video = list(wb["video"].values)
    assert video[0][0] == "video_id"
    assert len(video) == 3
    kolom = list(video[0])
    assert video[1][kolom.index("product_ids")] == "9,8"
    produk = list(wb["produk"].values)
    assert produk[1][1] == "Serum 😊"


def test_tulis_csv(tmp_path):
    paths = tulis_csv(tmp_path, "tokocontoh", *_data())
    assert [p.name for p in paths] == [
        "tokocontoh_profil.csv", "tokocontoh_video.csv", "tokocontoh_produk.csv",
    ]
    with paths[2].open(encoding="utf-8-sig") as f:
        baris = list(csv.DictReader(f))
    assert baris[0]["nama"] == "Serum 😊"
    assert baris[0]["komisi_persen"] == ""
